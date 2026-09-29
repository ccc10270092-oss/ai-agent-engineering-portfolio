"""知识文件文本提取：txt / md / csv / pdf / docx → 纯文本（供入库）。

- 纯文本：UTF-8（含 BOM）/ GBK 自动识别（中文 Windows 记事本常态）；
- CSV：表头嗅探 FAQ 问答列（问题/回答、question/answer、q/a）→ 「问/答」
  格式化，普通表格则按行拼接；
- PDF：pypdf 逐页提取（扫描版 PDF 无文本层 → 报错提示）；
- docx：python-docx 段落 + 表格。

解析库未安装时抛 ValueError（而非 ImportError），API 层统一转 400 反馈前端。
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

SUPPORTED_SUFFIXES = {".txt", ".md", ".csv", ".pdf", ".docx"}

_Q_HEADERS = {"问题", "question", "q", "问"}
_A_HEADERS = {"回答", "答案", "answer", "a", "答"}


def extract_text(data: bytes, filename: str) -> str:
    """从上传文件字节中提取纯文本；任何无法解析的情况抛 ValueError。"""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"不支持的文件类型 {suffix or '（无后缀）'}，"
            f"仅支持 {' / '.join(sorted(SUPPORTED_SUFFIXES))}"
        )
    if not data:
        raise ValueError("文件为空")

    if suffix in (".txt", ".md"):
        text = _decode_text(data)
    elif suffix == ".csv":
        text = _parse_csv(_decode_text(data))
    elif suffix == ".pdf":
        text = _parse_pdf(data)
    else:  # .docx
        text = _parse_docx(data)

    text = (text or "").strip()
    if not text:
        raise ValueError("未能从文件中提取到文本（扫描版 PDF 或空文档不支持）")
    return text


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码（支持 UTF-8 / GBK）")


def _parse_csv(text: str) -> str:
    reader = csv.DictReader(io.StringIO(text))
    headers = [h.strip().lower() for h in (reader.fieldnames or [])]
    if not headers:
        return text
    q_col = next((h for h in headers if h in _Q_HEADERS), None)
    a_col = next((h for h in headers if h in _A_HEADERS), None)

    if q_col and a_col:  # FAQ 问答表 → 逐条格式化，检索命中率高
        pairs = []
        for row in reader:
            q = (row.get(q_col) or "").strip()
            a = (row.get(a_col) or "").strip()
            if q or a:
                pairs.append(f"问：{q}\n答：{a}")
        return "\n\n".join(pairs)

    # 普通表格：每行非空单元格拼接为一行文本
    lines = []
    for row in reader:
        cells = [str(v).strip() for v in row.values() if v and str(v).strip()]
        if cells:
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("未安装 pypdf，无法解析 PDF（pip install pypdf）") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 - 损坏文件等统一转可读错误
        raise ValueError(f"PDF 解析失败：{exc}") from exc


def _parse_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ValueError("未安装 python-docx，无法解析 Word（pip install python-docx）") from exc
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Word 解析失败：{exc}") from exc
